"""Approximate slide renderer for layout QA (LibreOffice is unavailable here).

Draws each shape's box, fill and text with PIL at the same geometry PowerPoint
will use, so overlaps, overflow and edge violations are visible. Font metrics
are DejaVu rather than Calibri, so treat text width as indicative (±10%).
"""

import sys
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

DPI = 110
FONT_DIR = "/usr/share/fonts/truetype/dejavu"


def font(size_pt, bold=False):
    path = f"{FONT_DIR}/DejaVuSans{'-Bold' if bold else ''}.ttf"
    return ImageFont.truetype(path, max(int(size_pt * DPI / 72), 6))


def emu_px(value):
    return int(Emu(value).inches * DPI) if value is not None else 0


def rgb(colour, default=(20, 36, 59)):
    try:
        if colour and colour.type is not None and colour.rgb is not None:
            return tuple(bytes.fromhex(str(colour.rgb)))
    except Exception:
        pass
    return default


def wrap(draw, text, f, width):
    lines, current = [], ""
    for word in text.split():
        trial = (current + " " + word).strip()
        if draw.textlength(trial, font=f) <= width or not current:
            current = trial
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def render(path, prefix):
    prs = Presentation(path)
    W = int(Emu(prs.slide_width).inches * DPI)
    H = int(Emu(prs.slide_height).inches * DPI)
    outputs = []
    for index, slide in enumerate(prs.slides, start=1):
        img = Image.new("RGB", (W, H), "white")
        d = ImageDraw.Draw(img)
        for shape in slide.shapes:
            x, y = emu_px(shape.left), emu_px(shape.top)
            w, h = emu_px(shape.width), emu_px(shape.height)
            if shape.shape_type == 13:  # picture
                try:
                    pic = Image.open(io_bytes(shape)).convert("RGB")
                    img.paste(pic.resize((max(w, 1), max(h, 1))), (x, y))
                except Exception:
                    d.rectangle([x, y, x + w, y + h], outline=(150, 170, 190))
                continue
            fill = None
            try:
                if shape.fill.type is not None and shape.fill.type == 1:
                    fill = rgb(shape.fill.fore_color, (240, 244, 248))
            except Exception:
                fill = None
            if fill:
                d.rounded_rectangle([x, y, x + w, y + h], radius=6, fill=fill,
                                    outline=(210, 222, 234))
            if not shape.has_text_frame:
                continue
            ty = y + 2
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs) or para.text
                if not text.strip():
                    ty += 6
                    continue
                run = para.runs[0] if para.runs else None
                size = run.font.size.pt if (run is not None and run.font.size) else 12
                bold = bool(run is not None and run.font.bold)
                colour = rgb(run.font.color, (20, 36, 59)) if run is not None else (20, 36, 59)
                f = font(size, bold)
                for line in wrap(d, text, f, max(w - 6, 20)):
                    d.text((x + 3, ty), line, font=f, fill=colour)
                    ty += int(size * DPI / 72 * 1.18)
                ty += 2
            if ty > y + h + 4:                       # overflow marker
                d.rectangle([x, y, x + w, y + h], outline=(220, 30, 60), width=2)
                d.text((x + 2, y + h + 1), "OVERFLOW", font=font(7, True), fill=(220, 30, 60))
        d.rectangle([int(0.5 * DPI), int(0.5 * DPI), W - int(0.5 * DPI), H - int(0.5 * DPI)],
                    outline=(230, 235, 240))
        name = f"{prefix}-{index}.png"
        img.save(name)
        outputs.append(name)
    print("\n".join(outputs))


def io_bytes(shape):
    import io
    return io.BytesIO(shape.image.blob)


if __name__ == "__main__":
    render(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "render")
