"""Build the SIH 2026 Idea Presentation for PS 26165 inside the official template."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SRC, OUT = "sih_format.pptx", "SIH2026_WellDrop_SENTRA_PS26165.pptx"

NAVY = RGBColor(0x10, 0x24, 0x3B)
BLUE = RGBColor(0x00, 0x70, 0xC0)
RED = RGBColor(0xD7, 0x26, 0x3D)
AMBER = RGBColor(0xE0, 0x8A, 0x00)
GREEN = RGBColor(0x1E, 0x9E, 0x5A)
INK = RGBColor(0x10, 0x24, 0x3B)
MUTED = RGBColor(0x5A, 0x6B, 0x7C)
CARD = RGBColor(0xF2, 0xF6, 0xFA)
CARD_LINE = RGBColor(0xD5, 0xE2, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"


# --------------------------------------------------------------------- helpers
def textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    return frame


def write(frame, lines, *, first=True):
    """lines: (text, size, bold, colour, space_after, indent_level)."""
    for index, spec in enumerate(lines):
        text, size, bold, colour = spec[0], spec[1], spec[2], spec[3]
        space_after = spec[4] if len(spec) > 4 else 2
        level = spec[5] if len(spec) > 5 else 0
        para = frame.paragraphs[0] if (index == 0 and first) else frame.add_paragraph()
        para.level = level
        para.space_after = Pt(space_after)
        run = para.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = colour
        run.font.name = FONT
    return frame


def card(slide, x, y, w, h, *, fill=CARD, line=CARD_LINE, radius=0.06):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    shape.adjustments[0] = radius
    shape.text_frame.word_wrap = True
    return shape


def chip(slide, x, y, size, label, colour):
    """A small filled circle carrying a number or glyph."""
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                                   Inches(size), Inches(size))
    shape.fill.solid()
    shape.fill.fore_color.rgb = colour
    shape.line.fill.background()
    shape.shadow.inherit = False
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = label
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = FONT
    return shape


def bullet_rows(slide, x, y, w, rows, *, row_h=0.52, gap=0.06, colour=BLUE,
                title_size=10.5, body_size=9.5):
    """Stacked one-line cards: a coloured chip, a bold lead-in, then detail."""
    top = y
    for index, (lead, detail) in enumerate(rows, start=1):
        card(slide, x, top, w, row_h, fill=CARD, line=CARD_LINE)
        chip(slide, x + 0.12, top + (row_h - 0.26) / 2, 0.26, str(index), colour)
        frame = textbox(slide, x + 0.5, top + 0.05, w - 0.62, row_h - 0.1,
                        anchor=MSO_ANCHOR.MIDDLE)
        para = frame.paragraphs[0]
        para.space_after = Pt(0)
        for text, bold, size, tone in ((lead + "  ", True, title_size, INK),
                                       (detail, False, body_size, MUTED)):
            run = para.add_run()
            run.text = text
            run.font.bold = bold
            run.font.size = Pt(size)
            run.font.color.rgb = tone
            run.font.name = FONT
        top += row_h + gap
    return top


def section(slide, x, y, w, text, colour=BLUE, size=11):
    frame = textbox(slide, x, y, w, 0.24)
    write(frame, [(text, size, True, colour, 0)])
    return y + 0.28


def drop_shape(shape):
    shape._element.getparent().remove(shape._element)


def find(slide, *names):
    return [s for s in slide.shapes if s.name in names]


# ----------------------------------------------------------------------- build
prs = Presentation(SRC)
slides = prs.slides

# Remove the "important instructions" slide the template tells teams to delete.
sldIdLst = prs.slides._sldIdLst
last = list(sldIdLst)[-1]
prs.part.drop_rel(last.rId)
sldIdLst.remove(last)

# ------------------------------------------------------------------- slide 1
s1 = slides[0]
for shape in find(s1, "TextBox 9"):
    drop_shape(shape)
for shape in find(s1, "Subtitle 3"):
    frame = shape.text_frame
    frame.clear()
    write(frame, [("SENTRA", 40, True, NAVY, 0)])
    para = frame.add_paragraph()
    para.space_after = Pt(0)
    run = para.add_run()
    run.text = "Sense the Risk  ·  Stop the Incident"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = BLUE
    run.font.name = FONT

frame = textbox(s1, 0.42, 2.75, 6.6, 3.9)
write(frame, [
    ("Problem Statement ID – 26165", 14, True, INK, 7),
    ("Problem Statement Title – AI/NLP Engine to Detect Serious Injury & Fatality "
     "(SIF) Precursors in OIL's Unsafe-Act / Unsafe-Condition and Near-Miss Reports",
     12.5, False, INK, 7),
    ("Theme – Smart Automation", 12.5, False, INK, 7),
    ("PS Category – Software", 12.5, False, INK, 7),
    ("Team ID – ______________  (enter from the SIH portal)", 12.5, False, INK, 7),
    ("Team Name – WellDrop", 14, True, INK, 7),
    ("Organisation – Oil India Limited", 11.5, False, MUTED, 0),
])

# The template's team-name badge appears on every content slide.
for slide in slides:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if run.text.strip() == "Your Team Name":
                    run.text = "WellDrop"
                    run.font.size = Pt(11)
                    run.font.bold = True

# ------------------------------------------------------------------- slide 2
s2 = slides[1]
for shape in find(s2, "TextBox 8"):
    drop_shape(shape)

frame = textbox(s2, 0.45, 1.34, 12.45, 0.36)
write(frame, [
    ("SENTRA reads every UA/UC and near-miss report and flags fatal potential — "
     "on the shift it is written.", 14, True, INK, 0)])

left_w = 7.35
s2.shapes.add_picture("assets/rule.png", Inches(0.45), Inches(1.80),
                      width=Inches(left_w))
y = section(s2, 0.45, 3.70, left_w, "WHAT MAKES IT DIFFERENT")
bullet_rows(s2, 0.45, y, left_w, [
    ("Hybrid, not black box —",
     "safety rules + semantic transformer + XGBoost; the model may only add a flag, never overturn a rule"),
    ("Evidence on every flag —",
     "matched cues and decision path, so HSE audits the reasoning before acting"),
    ("Zero-shot rule mapping —",
     "9 IOGP Life-Saving Rules + 2 upstream categories, with no labelled data to start"),
    ("Runs offline —",
     "deterministic mode needs no model and no network: safe for plant/air-gapped sites"),
    ("Human-in-the-loop —",
     "rule/model disagreement is queued for an expert, and that decision becomes a label"),
], row_h=0.5, gap=0.055)

right_x, right_w = 8.10, 4.80
c = card(s2, right_x, 1.80, right_w, 1.32, fill=RGBColor(0xFD, 0xF1, 0xF2),
         line=RGBColor(0xF2, 0xC7, 0xCC))
frame = textbox(s2, right_x + 0.18, 1.92, right_w - 0.36, 1.1)
write(frame, [
    ("THE GAP TODAY", 10.5, True, RED, 5),
    ("Reports triaged by hand, monthly or quarterly", 10, False, INK, 3),
    ("Injuries fell 51% in 15 years — fatalities only 25.5%", 10, False, INK, 3),
    ("Low severity ≠ low fatal potential: ~20–25% of reports carry it", 10, False, INK, 0),
])

s2.shapes.add_picture("assets/dash_top.png", Inches(right_x), Inches(3.30),
                      width=Inches(right_w))
frame = textbox(s2, right_x, 5.22, right_w, 0.3)
write(frame, [("Working prototype — live dashboard, not a mock-up", 9, False, MUTED, 0)])

stats = [("74", "automated tests"), ("11", "IOGP categories"), ("0", "labelled rows to start")]
sx = right_x
for value, label in stats:
    card(s2, sx, 5.72, 1.5, 0.78, fill=CARD, line=CARD_LINE)
    frame = textbox(s2, sx, 5.80, 1.5, 0.64)
    para = frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    para.space_after = Pt(0)
    run = para.add_run()
    run.text = value
    run.font.size = Pt(19)
    run.font.bold = True
    run.font.color.rgb = BLUE
    run.font.name = FONT
    para2 = frame.add_paragraph()
    para2.alignment = PP_ALIGN.CENTER
    run2 = para2.add_run()
    run2.text = label
    run2.font.size = Pt(8.5)
    run2.font.color.rgb = MUTED
    run2.font.name = FONT
    sx += 1.65

# ------------------------------------------------------------------- slide 3
s3 = slides[2]
for shape in find(s3, "TextBox 8"):
    drop_shape(shape)

s3.shapes.add_picture("assets/pipeline.png", Inches(0.45), Inches(1.28),
                      width=Inches(12.45))

y = 3.86
col_w = 6.0
section(s3, 0.45, y, col_w, "TECHNOLOGY STACK")
stack = [
    ("Language & UI", "Python 3.11 · PyQt6 desktop console · QThread workers"),
    ("NLP", "sentence-transformers all-MiniLM-L6-v2 · rule + prototype knowledge base"),
    ("ML / MLOps", "XGBoost · scikit-learn · MLflow experiment tracking"),
    ("Ingestion", "PaddleOCR · pypdfium2 PDF text layer · CSV importer"),
    ("Engineering", "NumPy · unittest (74 tests) · Git / GitHub"),
]
top = y + 0.3
for label, detail in stack:
    frame = textbox(s3, 0.45, top, col_w, 0.42)
    para = frame.paragraphs[0]
    para.space_after = Pt(0)
    for text, bold, size, tone in ((label + " — ", True, 10, BLUE),
                                   (detail, False, 9.5, INK)):
        run = para.add_run()
        run.text = text
        run.font.bold = bold
        run.font.size = Pt(size)
        run.font.color.rgb = tone
        run.font.name = FONT
    top += 0.44

rx, rw = 6.85, 6.05
section(s3, rx, y, rw, "METHODOLOGY")
bullet_rows(s3, rx, y + 0.3, rw, [
    ("Rules first —", "auditable deterministic backbone"),
    ("Semantics second —", "zero-shot prototypes lift recall on paraphrase"),
    ("Model third —", "XGBoost opinion, MLflow-tracked, never an override"),
    ("Loop closed —", "expert review produces the labels for retraining"),
], row_h=0.44, gap=0.05, colour=NAVY)

card(s3, rx, 6.12, rw, 0.62, fill=RGBColor(0xEC, 0xF7, 0xF1),
     line=RGBColor(0xBF, 0xE4, 0xD3))
frame = textbox(s3, rx + 0.16, 6.20, rw - 0.32, 0.5)
write(frame, [
    ("Guardrail: ", 10, True, GREEN, 0),
])
para = frame.paragraphs[0]
run = para.add_run()
run.text = ("if an encoder scores every prototype alike, its ranking is ignored and the "
            "rules decide — verified against a deliberately degenerate model.")
run.font.size = Pt(9.5)
run.font.color.rgb = INK
run.font.name = FONT

# ------------------------------------------------------------------- slide 4
s4 = slides[3]
for shape in find(s4, "TextBox 8"):
    drop_shape(shape)

left_w = 4.35
y = section(s4, 0.45, 1.30, left_w, "WHY IT IS FEASIBLE", GREEN)
frame = textbox(s4, 0.45, y, left_w, 1.85)
write(frame, [
    ("•  Laptop-class hardware — no GPU, no server", 10, False, INK, 5),
    ("•  Offline mode needs no model and no network", 10, False, INK, 5),
    ("•  Reads existing HSSE exports (CSV, PDF, scans) — crews keep reporting as they do",
     10, False, INK, 5),
    ("•  Open-source stack, deployed site by site", 10, False, INK, 5),
    ("•  Already built and tested — 74 automated tests pass", 10, False, INK, 0),
])

y2 = section(s4, 0.45, 3.20, left_w, "VIABILITY", BLUE)
frame = textbox(s4, 0.45, y2, left_w, 1.85)
write(frame, [
    ("•  Low cost — open-source components, no per-seat licence", 10, False, INK, 5),
    ("•  Return from HSE screening time saved and incidents prevented", 10, False, INK, 5),
    ("•  One engine serves every asset; scales site by site", 10, False, INK, 5),
    ("•  Improves with use — each expert decision becomes a label", 10, False, INK, 5),
    ("•  Supports OIL's statutory reporting rather than replacing it", 10, False, INK, 0),
])

card(s4, 0.45, 5.12, left_w, 1.62, fill=RGBColor(0xEC, 0xF3, 0xFA),
     line=RGBColor(0xC3, 0xDC, 0xF0))
frame = textbox(s4, 0.62, 5.24, left_w - 0.34, 1.4)
write(frame, [
    ("WHAT A PILOT LOOKS LIKE", 10, True, BLUE, 5),
    ("Week 0 — point it at one asset's existing report export", 9.5, False, INK, 4),
    ("Weeks 1–8 — HSE works the review queue; decisions become labels", 9.5, False, INK, 4),
    ("Week 9 — first model trained on real labels, promoted only at the recall gate",
     9.5, False, INK, 0),
])

rx, rw = 5.15, 7.75
section(s4, rx, 1.30, rw, "RISKS AND HOW WE HANDLE THEM", AMBER)
risks = [
    ("No labelled SIF corpus exists",
     "Start rule-driven; review queue produces labels; promote a model only at ≥200 labels and recall ≥0.80"),
    ("A missed precursor (false negative)",
     "Semantic layer may only add flags; rules are never overridden; recall-first acceptance gate"),
    ("Plant networks are air-gapped",
     "Pre-fetch the encoder and OCR models; full offline lexical mode as fallback"),
    ("Small samples skew site rankings",
     "Density ranked by Wilson lower bound — 2-of-2 cannot outrank 20-of-30"),
    ("Scanned and handwritten reports",
     "PDF text layer first, PaddleOCR for scans, per-line confidence routed to review"),
    ("Model drift after go-live",
     "MLflow run history plus monitored KPIs: agreement %, unclassified share, queue age"),
]
top = 1.62
for index, (risk, fix) in enumerate(risks, start=1):
    card(s4, rx, top, rw, 0.78, fill=CARD, line=CARD_LINE)
    chip(s4, rx + 0.14, top + 0.26, 0.26, str(index), AMBER)
    frame = textbox(s4, rx + 0.52, top + 0.09, rw - 0.68, 0.62)
    write(frame, [(risk, 10, True, INK, 2)])
    para = frame.add_paragraph()
    para.space_after = Pt(0)
    run = para.add_run()
    run.text = fix
    run.font.size = Pt(9.5)
    run.font.color.rgb = MUTED
    run.font.name = FONT
    top += 0.855

# ------------------------------------------------------------------- slide 5
s5 = slides[4]
for shape in find(s5, "TextBox 8"):
    drop_shape(shape)

headline = [
    ("Monthly", "manual triage today", RED),
    ("Same shift", "with SENTRA", GREEN),
    ("20–25%", "of reports hold the fatal potential", BLUE),
    ("Every flag", "carries its evidence", NAVY),
]
hx = 0.45
for value, label, tone in headline:
    card(s5, hx, 1.30, 3.02, 0.86, fill=CARD, line=CARD_LINE)
    frame = textbox(s5, hx + 0.16, 1.40, 2.7, 0.7)
    write(frame, [(value, 17, True, tone, 1), (label, 9.5, False, MUTED, 0)])
    hx += 3.16

section(s5, 0.45, 2.38, 12.45, "WHO BENEFITS, AND HOW")
groups = [
    ("HSE teams & asset owners", BLUE, [
        "Screening effort moves from every report to the ones that can kill",
        "Sites and activities ranked by precursor density, not by report volume",
        "Recurring failed barriers named — fix the control, not the incident",
    ]),
    ("Workers & contractors", GREEN, [
        "Exposure removed before the event, not investigated after it",
        "Reporting visibly acted on, which strengthens reporting culture",
        "Same protection for contractor crews, who carry much of the field risk",
    ]),
    ("Economic, social & environmental", AMBER, [
        "Avoided incident cost, downtime and emergency response",
        "Audit-ready evidence trail for regulators and internal review",
        "Paperless triage; less duplicated manual review each cycle",
    ]),
]
gx = 0.45
for title, tone, points in groups:
    card(s5, gx, 2.72, 4.05, 1.98, fill=CARD, line=CARD_LINE)
    frame = textbox(s5, gx + 0.18, 2.86, 3.7, 2.1)
    write(frame, [(title, 11, True, tone, 6)])
    for point in points:
        para = frame.add_paragraph()
        para.space_after = Pt(5)
        run = para.add_run()
        run.text = "•  " + point
        run.font.size = Pt(9.5)
        run.font.color.rgb = INK
        run.font.name = FONT
    gx += 4.24

s5.shapes.add_picture("assets/review.png", Inches(0.45), Inches(4.88), width=Inches(6.30))
frame = textbox(s5, 7.05, 4.94, 5.85, 1.6)
write(frame, [
    ("THE HUMAN STAYS IN CHARGE", 10.5, True, NAVY, 6),
    ("Disagreements, critical risk and thin evidence are queued with the reason "
     "attached. No report is closed, downgraded or escalated by the engine.",
     9.5, False, INK, 6),
    ("Each expert decision returns as a training label.", 9.5, True, GREEN, 0),
])

# ------------------------------------------------------------------- slide 6
s6 = slides[5]
for shape in find(s6, "TextBox 8"):
    drop_shape(shape)

y = section(s6, 0.45, 1.30, 7.4, "PEER-REVIEWED RESEARCH")
papers = [
    ("Parikh et al. (2024), Scientific Reports",
     "Automatic identification of incidents involving potential serious injuries and "
     "fatalities — Transformer + XGBoost, recall-first.  nature.com/articles/s41598-024-58824-y"),
    ("Ganguli et al. (2021), Minerals",
     "NLP-based machine learning on mine incident narratives.  mdpi.com/2075-163X/11/7/776"),
    ("Chang & Martinez (2024), Process Safety and Environmental Protection",
     "NLP for spill reduction in an E&P company — causes and contributing factors."),
    ("Wang & Wang (2026), J. Loss Prevention in the Process Industries",
     "Bow-tie grounded sentence-level annotation for interpretable incident analysis."),
]
top = y
for index, (cite, detail) in enumerate(papers, start=1):
    card(s6, 0.45, top, 7.4, 0.86, fill=CARD, line=CARD_LINE)
    chip(s6, 0.59, top + 0.30, 0.26, str(index), BLUE)
    frame = textbox(s6, 0.97, top + 0.11, 6.7, 0.7)
    write(frame, [(cite, 10, True, INK, 2)])
    para = frame.add_paragraph()
    para.space_after = Pt(0)
    run = para.add_run()
    run.text = detail
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    run.font.name = FONT
    top += 0.94

card(s6, 0.45, 5.30, 7.4, 1.42, fill=RGBColor(0xEC, 0xF7, 0xF1),
     line=RGBColor(0xBF, 0xE4, 0xD3))
frame = textbox(s6, 0.62, 5.42, 7.06, 1.2)
write(frame, [
    ("HOW THE RESEARCH SHAPED SENTRA", 10, True, GREEN, 5),
    ("[1] recall-first evaluation and the transformer + XGBoost pairing", 9.5, False, INK, 4),
    ("[2][3] evidence that NLP generalises across mining and E&P narratives", 9.5, False, INK, 4),
    ("[4] barrier-level annotation — why we classify on energy and barrier, not keywords",
     9.5, False, INK, 0),
])

rx, rw = 8.15, 4.75
y = section(s6, rx, 1.30, rw, "STANDARDS & INDUSTRY MODELS")
frame = textbox(s6, rx, y, rw, 1.5)
write(frame, [
    ("•  IOGP Life-Saving Rules (2018) — the 9-rule taxonomy we map to", 9.5, False, INK, 5),
    ("•  EEI SIF precursor model — energy + failed barrier logic", 9.5, False, INK, 5),
    ("•  DEKRA, Martin & Black (2015) — SIF potential vs actual severity", 9.5, False, INK, 5),
    ("•  VelocityEHS PSIF classifier (2024) — commercial benchmark", 9.5, False, INK, 0),
])

y = section(s6, rx, 3.05, rw, "HOW WE DIFFER FROM EXISTING TOOLS")
frame = textbox(s6, rx, y, rw, 2.0)
write(frame, [
    ("•  Commercial PSIF tools need a large labelled corpus; SENTRA starts with none",
     9.5, False, INK, 5),
    ("•  Most classifiers output a score; SENTRA outputs the cues and the decision path",
     9.5, False, INK, 5),
    ("•  Built for the plant — full offline mode, no cloud dependency", 9.5, False, INK, 5),
    ("•  Ranks by precursor density, so a small busy site is not hidden by a large one",
     9.5, False, INK, 0),
])

card(s6, rx, 5.55, rw, 0.95, fill=RGBColor(0xEC, 0xF3, 0xFA), line=RGBColor(0xC3, 0xDC, 0xF0))
frame = textbox(s6, rx + 0.18, 5.68, rw - 0.36, 0.75)
write(frame, [
    ("PROTOTYPE SOURCE", 10, True, BLUE, 4),
    ("github.com/tedo001/SIF — working code, 74 automated tests, operating manual",
     9.5, False, INK, 0),
])

# ---------------------------------------------------------------- speaker notes
NOTES = {
    0: ("Open with the number: OIL already collects the reports. The question is which "
        "of them could have killed someone. Team WellDrop; the engine is SENTRA."),
    1: ("The one idea to land: SIF potential is high energy AND a failed barrier - not "
        "keyword matching, and not the severity that actually occurred. Everything else "
        "follows from that rule. The screenshot is our running prototype, not a mock-up."),
    2: ("Walk the seven stages left to right in one sentence each. Emphasise the green "
        "return line: expert decisions come back as labels, so the engine improves with "
        "use. If asked why not an LLM: auditability - every flag names the cues that "
        "produced it."),
    3: ("Lead with the honest risk - no labelled SIF corpus exists anywhere - then the "
        "answer: the rules work from day one and the review queue manufactures the "
        "labels. Offline mode is what makes it deployable on a plant network."),
    4: ("Impact is a reallocation of attention, not a new headcount: the same HSE team "
        "spends its hours on the fifth of reports that carry fatal potential. Stress that "
        "no report is ever closed by the engine."),
    5: ("If challenged on novelty: commercial PSIF tools need a labelled corpus and return "
        "a score; we start with none and return the evidence. Offer the repository - the "
        "code, the tests and the operating manual are public."),
}
for index, note in NOTES.items():
    slides[index].notes_slide.notes_text_frame.text = note

prs.save(OUT)
print("written", OUT, "with", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
